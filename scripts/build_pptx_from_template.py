# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx"]
# ///
"""Genera un .pptx nuevo partiendo de una plantilla.

Dos formas de definir cada diapositiva:
  * "layout": crea una diapositiva vacía con ese layout (nombre o índice) y rellena placeholders por idx.
  * "from_slide": CLONA una diapositiva de ejemplo de la plantilla (nº 1-based). Es lo adecuado cuando la
    plantilla es una galería de diseños, porque conserva imágenes, formas y decoración.

Uso: build_pptx_from_template.py <plantilla.pptx> <salida.pptx> <spec.json>

spec.json:
{ "slides": [
    {"from_slide": 4, "texts": {"1": "Título", "2": "Subtítulo"}},
    {"from_slide": 50, "texts": {"1": "Mensaje\\nclave", "2": ["Párrafo 1", "Párrafo 2"]}, "delete": [3]},
    {"from_slide": 53, "table": [["", "2026", "2027"], ["Ventas", "1", "2"]]},
    {"layout": "Blank"},
    {"layout": 1, "placeholders": {"0": "Agenda", "1": ["Uno", "Dos"]}, "notes": "Notas"}
] }

"texts" usa el nº de orden de las formas con texto de la diapositiva, tal como lo muestra
`list_styles.py <pptx> --slides`. Se conserva el formato del primer fragmento de cada forma.
"delete" elimina formas por ese mismo nº (por ejemplo, ítems de índice que no uses).
"table" rellena la primera tabla (celdas en orden; añade o quita filas por clonado si hace falta).
"""
import copy
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def find_layout(prs, ref):
    if isinstance(ref, int):
        return prs.slide_layouts[ref]
    for lay in prs.slide_layouts:
        if lay.name.lower() == str(ref).lower():
            return lay
    raise SystemExit(f"Layout inexistente: {ref!r}. Disponibles: {[l.name for l in prs.slide_layouts]}")


def text_shapes(slide):
    """Formas con texto en orden de documento (misma numeración que list_styles --slides)."""
    return [sh for sh in slide.shapes if sh.has_text_frame and not sh.shape_type == 19]


def set_text(shape, value):
    lines = value if isinstance(value, list) else str(value).split("\n")
    tf = shape.text_frame
    paras = tf.paragraphs
    tmpl = copy.deepcopy(paras[0]._p) if paras else None
    # deja un solo párrafo con un solo run (formato del primero)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = lines[0]
    else:
        p0.text = lines[0]
    for br in p0._p.findall(qn("a:br")):
        p0._p.remove(br)
    last = p0._p
    for line in lines[1:]:
        newp = copy.deepcopy(tmpl)
        last.addnext(newp)
        last = newp
        rs = newp.findall(qn("a:r"))
        for r in rs[1:]:
            newp.remove(r)
        for br in newp.findall(qn("a:br")):
            newp.remove(br)
        if rs:
            rs[0].find(qn("a:t")).text = line
        else:
            tf.paragraphs[-1].text = line


def clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        c = copy.deepcopy(el)
        for node in c.iter():
            for attr in list(node.attrib):
                if attr.startswith("{" + R_NS + "}") and node.attrib[attr] in rid_map:
                    node.set(attr, rid_map[node.attrib[attr]])
        new.shapes._spTree.append(c)
    bg = src._element.find(qn("p:cSld")).find(qn("p:bg"))
    if bg is not None:
        new._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))
    return new


def fill_table(slide, rows):
    tbl = next((sh.table for sh in slide.shapes if sh.has_table), None)
    if tbl is None:
        raise SystemExit("La diapositiva clonada no tiene tabla")
    trs = tbl._tbl.findall(qn("a:tr"))
    while len(trs) < len(rows):  # clona la última fila
        trs[-1].addnext(copy.deepcopy(trs[-1]))
        trs = tbl._tbl.findall(qn("a:tr"))
    for tr in trs[len(rows):]:
        tbl._tbl.remove(tr)
    for r, row in enumerate(rows):
        for c, val in enumerate(row[: len(tbl.columns)]):
            cell = tbl.cell(r, c)
            paras = cell.text_frame.paragraphs
            if paras[0].runs:
                paras[0].runs[0].text = str(val)
                for rn in paras[0].runs[1:]:
                    rn._r.getparent().remove(rn._r)
            else:
                paras[0].text = str(val)


def main() -> int:
    src, dst, spec_path = sys.argv[1:4]
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8-sig"))
    prs = Presentation(src)
    originals = list(prs.slides)
    n_orig = len(originals)

    for n, s in enumerate(spec["slides"], 1):
        if "from_slide" in s:
            i = int(s["from_slide"])
            if not 1 <= i <= n_orig:
                raise SystemExit(f"Diapositiva {n}: from_slide={i} fuera de rango (1..{n_orig})")
            slide = clone_slide(prs, originals[i - 1])
            shapes = text_shapes(slide)
            for k, v in s.get("texts", {}).items():
                if not 1 <= int(k) <= len(shapes):
                    raise SystemExit(f"Diapositiva {n}: no existe la forma de texto nº {k} (hay {len(shapes)})")
                set_text(shapes[int(k) - 1], v)
            for k in sorted((int(x) for x in s.get("delete", [])), reverse=True):
                el = shapes[k - 1]._element
                el.getparent().remove(el)
            if "table" in s:
                fill_table(slide, s["table"])
        else:
            slide = prs.slides.add_slide(find_layout(prs, s["layout"]))
            by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
            for idx, value in s.get("placeholders", {}).items():
                ph = by_idx.get(int(idx))
                if ph is None:
                    raise SystemExit(f"Diapositiva {n}: el layout no tiene placeholder idx={idx}. Hay: {sorted(by_idx)}")
                set_text(ph, value)
            for ph in list(slide.placeholders):
                if ph.has_text_frame and not ph.text_frame.text.strip():
                    ph._element.getparent().remove(ph._element)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    # elimina las diapositivas de ejemplo originales
    sld_ids = prs.slides._sldIdLst
    for sld in list(sld_ids)[:n_orig]:
        prs.part.drop_rel(sld.rId)
        sld_ids.remove(sld)
    prs.save(dst)
    print(f"OK -> {dst} ({len(spec['slides'])} diapositivas)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
