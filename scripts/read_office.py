# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "python-pptx", "openpyxl", "pymupdf"]
# ///
"""Lee docx/pptx/xlsx/pdf y lo vuelca como texto estructurado (estilo Markdown).

Uso: read_office.py <archivo> [--max-rows N]
- docx: párrafos con su estilo y tablas (incluye celdas dentro de controles de contenido).
- pptx: por diapositiva, layout, placeholders (idx) y notas.
- xlsx: por hoja, celdas con coordenadas, valor y fórmula.
- pdf : texto por página.
"""
import argparse
import sys
from pathlib import Path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def read_docx(path: Path, _max_rows: int) -> None:
    from docx import Document
    from docx.text.paragraph import Paragraph

    d = Document(str(path))
    t_idx = 0
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, d)
            if p.text.strip():
                print(f"[{p.style.name}] {p.text}")
        elif tag == "tbl":
            print(f"\n### Tabla {t_idx}")
            for tr in el.iter(f"{W}tr"):
                cells = ["".join(t.text or "" for t in tc.iter(f"{W}t")).replace("\n", " ") for tc in tr.iter(f"{W}tc")]
                print("| " + " | ".join(cells) + " |")
            print()
            t_idx += 1
        elif tag == "sdt":
            txt = "".join(t.text or "" for t in el.iter(f"{W}t"))
            if txt.strip():
                print(f"[control de contenido] {txt}")
    for i, s in enumerate(d.sections):
        h = " / ".join(p.text for p in s.header.paragraphs if p.text.strip())
        f = " / ".join(p.text for p in s.footer.paragraphs if p.text.strip())
        if h or f:
            print(f"\n(sección {i}) cabecera: {h!r} pie: {f!r}")


def read_pptx(path: Path, _max_rows: int) -> None:
    from pptx import Presentation

    p = Presentation(str(path))
    for n, s in enumerate(p.slides, 1):
        print(f"\n## Diapositiva {n} (layout: {s.slide_layout.name})")
        for sh in s.shapes:
            kind = f"placeholder idx={sh.placeholder_format.idx}" if sh.is_placeholder else sh.shape_type
            if sh.has_text_frame and sh.text_frame.text.strip():
                print(f"- [{sh.name} | {kind}] {sh.text_frame.text.strip()!r}")
            elif getattr(sh, "has_table", False) and sh.has_table:
                print(f"- [{sh.name} | tabla]")
                for r in sh.table.rows:
                    print("  | " + " | ".join(c.text for c in r.cells) + " |")
            elif sh.shape_type == 13:
                print(f"- [{sh.name} | imagen]")
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            print(f"  notas: {s.notes_slide.notes_text_frame.text.strip()!r}")


def read_xlsx(path: Path, max_rows: int) -> None:
    from openpyxl import load_workbook

    wf = load_workbook(str(path))
    wv = load_workbook(str(path), data_only=True)
    for ws in wf.worksheets:
        print(f"\n## Hoja {ws.title!r} dims={ws.dimensions} congelado={ws.freeze_panes}")
        vs = wv[ws.title]
        for r, row in enumerate(ws.iter_rows(), 1):
            if r > max_rows:
                print(f"... (más de {max_rows} filas)")
                break
            for c in row:
                if c.value is None:
                    continue
                if isinstance(c.value, str) and c.value.startswith("="):
                    print(f"{c.coordinate}: {c.value}  -> {vs[c.coordinate].value!r}")
                else:
                    print(f"{c.coordinate}: {c.value!r}")


def read_pdf(path: Path, _max_rows: int) -> None:
    import pymupdf

    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, 1):
            print(f"\n## Página {i}\n{page.get_text().strip()}")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf8")
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--max-rows", type=int, default=200)
    a = ap.parse_args()
    path = Path(a.file)
    fn = {".docx": read_docx, ".dotx": read_docx, ".pptx": read_pptx, ".potx": read_pptx,
          ".xlsx": read_xlsx, ".xlsm": read_xlsx, ".pdf": read_pdf}.get(path.suffix.lower())
    if not fn:
        print(f"Formato no soportado: {path.suffix}", file=sys.stderr)
        return 2
    fn(path, a.max_rows)
    return 0


if __name__ == "__main__":
    sys.exit(main())
