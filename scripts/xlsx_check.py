# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Auditoría de un .xlsx: fórmulas, errores calculados, fórmulas sin valor en caché, números 'a mano'.

Uso: xlsx_check.py <archivo.xlsx>
Los valores en caché solo existen si Excel ha calculado el libro: ejecuta antes
scripts/recalc_xlsx.ps1 (Windows) o recalc_xlsx.sh (macOS) si has generado fórmulas con openpyxl.
Sale con código 1 si hay errores (#REF!, #NAME?, #DIV/0!...).
"""
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf8")
ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
path = sys.argv[1]
wf = load_workbook(path)
wv = load_workbook(path, data_only=True)
bad = uncached = 0
for ws in wf.worksheets:
    f = h = 0
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                f += 1
                cached = wv[ws.title][c.coordinate].value
                if cached is None:
                    uncached += 1
                elif isinstance(cached, str) and cached in ERRORS:
                    bad += 1
                    print(f"ERROR {ws.title}!{c.coordinate}: {v} -> {cached}")
            elif isinstance(v, (int, float)):
                h += 1
    print(f"{ws.title}: fórmulas={f} números constantes={h} dims={ws.dimensions}")
if uncached:
    print(f"AVISO: {uncached} fórmulas sin valor calculado (recalcula con Excel).")
print("OK" if not bad else f"{bad} errores")
sys.exit(1 if bad else 0)
