# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "python-pptx", "openpyxl"]
# ///
"""Analiza una plantilla (.docx/.dotx/.pptx/.potx/.xlsx) y emite un perfil de marca en JSON.

Uso: extract_brand.py <plantilla> [-o brand-profile.json]
"""
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path


def theme_info(path: Path) -> dict:
    """Colores y fuentes del tema (comunes a los tres formatos OOXML)."""
    out = {"colors": {}, "fonts": {}}
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.search(r"theme/theme\d*\.xml$", n)]
        if not names:
            return out
        xml = z.read(sorted(names)[0]).decode("utf8", "ignore")
    for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6", "hlink", "folHlink"):
        m = re.search(rf"<a:{tag}>.*?(?:val|lastClr)=\"([0-9A-Fa-f]{{6}})\"", xml, re.S)
        if m:
            out["colors"][tag] = "#" + m.group(1).upper()
    for kind in ("majorFont", "minorFont"):
        m = re.search(rf"<a:{kind}>\s*<a:latin typeface=\"([^\"]*)\"", xml)
        if m:
            out["fonts"][kind] = m.group(1)
    return out


def media(path: Path) -> list:
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if "/media/" in n]


def docx_info(path: Path) -> dict:
    from docx import Document
    d = Document(str(path))
    styles = []
    for s in d.styles:
        if s.type == 1 and (s.builtin is False or s.name.startswith(("Heading", "Title", "Normal", "List", "Caption"))):
            f = s.font
            styles.append({
                "name": s.name,
                "font": f.name,
                "size_pt": f.size.pt if f.size else None,
                "bold": f.bold,
                "color": str(f.color.rgb) if f.color is not None and f.color.type is not None and f.color.rgb else None,
            })
    sec = d.sections[0]
    return {
        "paragraph_styles": styles,
        "page": {
            "width_cm": round(sec.page_width.cm, 2), "height_cm": round(sec.page_height.cm, 2),
            "margins_cm": {k: round(getattr(sec, f"{k}_margin").cm, 2) for k in ("top", "bottom", "left", "right")},
        },
        "header_text": [p.text for p in sec.header.paragraphs if p.text.strip()],
        "footer_text": [p.text for p in sec.footer.paragraphs if p.text.strip()],
        "table_styles": [s.name for s in d.styles if s.type == 3],
        "body_paragraphs": len(d.paragraphs),
    }


def pptx_info(path: Path) -> dict:
    from pptx import Presentation
    p = Presentation(str(path))
    layouts = []
    for i, l in enumerate(p.slide_layouts):
        layouts.append({
            "index": i, "name": l.name,
            "placeholders": [{"idx": ph.placeholder_format.idx, "type": str(ph.placeholder_format.type), "name": ph.name}
                             for ph in l.placeholders],
        })
    return {
        "slide_size_cm": [round(p.slide_width / 360000, 2), round(p.slide_height / 360000, 2)],
        "layouts": layouts, "existing_slides": len(p.slides),
    }


def xlsx_info(path: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    sheets = []
    for ws in wb.worksheets:
        c = ws["A1"]
        sheets.append({
            "name": ws.title, "dimensions": ws.dimensions,
            "A1_font": {"name": c.font.name, "size": c.font.sz, "bold": c.font.b},
            "A1_fill": c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None,
            "freeze_panes": ws.freeze_panes,
            "col_widths": {k: round(v.width, 1) for k, v in list(ws.column_dimensions.items())[:10] if v.width},
        })
    return {"sheets": sheets}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf8")
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    path = Path(a.template)
    ext = path.suffix.lower()
    kind = {".docx": "docx", ".dotx": "docx", ".pptx": "pptx", ".potx": "pptx", ".xlsx": "xlsx", ".xltx": "xlsx"}.get(ext)
    if not kind:
        print(f"Formato no soportado: {ext}", file=sys.stderr)
        return 2
    profile = {"source": path.name, "kind": kind, "theme": theme_info(path), "media": media(path),
               **{"docx": docx_info, "pptx": pptx_info, "xlsx": xlsx_info}[kind](path)}
    text = json.dumps(profile, indent=2, ensure_ascii=False)
    if a.out:
        Path(a.out).write_text(text, encoding="utf8")
    print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
