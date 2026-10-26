# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "python-pptx"]
# ///
"""Lista lo que una plantilla permite usar.

docx: estilos de párrafo/tabla/carácter, con fuente, tamaño, negrita y color.
pptx: layouts con sus placeholders (idx, tipo, nombre).

Uso: list_styles.py <plantilla.docx|.pptx>
"""
import sys
from pathlib import Path


def docx_styles(path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE as T

    d = Document(path)
    kinds = {T.PARAGRAPH: "párrafo", T.TABLE: "tabla", T.CHARACTER: "carácter"}
    for kind, label in kinds.items():
        print(f"## Estilos de {label}")
        for s in d.styles:
            if s.type != kind or s.hidden:
                continue
            f = getattr(s, "font", None)
            extra = ""
            if f is not None:
                color = f.color.rgb if f.color is not None and f.color.type is not None and f.color.rgb else None
                extra = f"  [{f.name or '-'} {f.size.pt if f.size else '-'}pt{' negrita' if f.bold else ''}{f' #{color}' if color else ''}]"
            print(f"- {s.name}{extra}")
        print()


def pptx_layouts(path):
    from pptx import Presentation

    p = Presentation(path)
    print(f"Tamaño: {p.slide_width.inches:.2f} x {p.slide_height.inches:.2f} in")
    for i, lay in enumerate(p.slide_layouts):
        print(f"[{i}] {lay.name}")
        for ph in lay.placeholders:
            f = ph.placeholder_format
            print(f"     idx={f.idx} tipo={f.type} nombre={ph.name!r}")


def pptx_slides(path):
    """Diapositivas de ejemplo con sus formas de texto numeradas (lo que usa build_pptx_from_template)."""
    from pptx import Presentation

    for i, s in enumerate(Presentation(path).slides, 1):
        print(f"## Diapositiva {i} (layout: {s.slide_layout.name})")
        k = 0
        for sh in s.shapes:
            if sh.has_text_frame and sh.shape_type != 19:
                k += 1
                t = sh.text_frame.text.replace("\n", " / ").replace("\x0b", " ")
                print(f"  {k}. {sh.name}: {t[:70]!r}")
            elif sh.has_table:
                print(f"  (tabla {len(sh.table.rows)}x{len(sh.table.columns)})")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    path = sys.argv[1]
    ext = Path(path).suffix.lower()
    if ext in (".pptx", ".potx") and "--slides" in sys.argv:
        return pptx_slides(path) or 0
    if ext in (".docx", ".dotx"):
        docx_styles(path)
    elif ext in (".pptx", ".potx"):
        pptx_layouts(path)
    else:
        raise SystemExit("Formato no soportado (usa docx o pptx)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
