# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "python-pptx", "openpyxl", "reportlab"]
# ///
"""Crea documentos de prueba en la carpeta indicada. Uso: make_fixtures.py <carpeta>"""
import sys
from pathlib import Path

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

from docx import Document
d = Document()
d.add_heading("Informe {{cliente}}", 1)
p = d.add_paragraph()
p.add_run("Hola ").bold = True
p.add_run("{{cl")
p.add_run("iente}} bienvenido")
t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "Año"; t.cell(0, 1).text = "2025"
d.sections[0].header.paragraphs[0].text = "Cabecera {{cliente}}"
d.save(out / "t.docx")

from pptx import Presentation
pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[1])
s.shapes.title.text = "Resumen {{cliente}}"
s.placeholders[1].text = "Cuerpo 2025"
pr.save(out / "t.pptx")

from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = "Datos"
ws["A1"], ws["A2"], ws["A3"] = 10, 20, "=SUM(A1:A2)"
ws["B1"] = "=1/0"
wb.save(out / "t.xlsx")

from reportlab.pdfgen import canvas
for n in (1, 2):
    c = canvas.Canvas(str(out / f"p{n}.pdf"))
    c.drawString(100, 750, f"Hola PDF {n}"); c.showPage(); c.drawString(100, 750, f"Pagina 2 de {n}"); c.save()
print("ok")
